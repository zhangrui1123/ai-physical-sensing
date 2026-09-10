# -*- coding: utf-8 -*-
"""TimesFM-3 主要想法介绍 PPT（依据 Google Research 博客 2026-08-31）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0x13, 0x1C, 0x2E)
CARD2 = RGBColor(0x18, 0x24, 0x3A)
STROKE = RGBColor(0x2A, 0x3A, 0x55)
CYAN = RGBColor(0x38, 0xBD, 0xF8)
INDIGO = RGBColor(0x81, 0x8C, 0xF8)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
GREEN = RGBColor(0x34, 0xD3, 0x99)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
SOFT = RGBColor(0xCB, 0xD5, 0xE1)
FONT = "Microsoft YaHei"
TOTAL = 8


def set_run(run, size, color, bold=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def put(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        set_run(run, size, color, bold)
    return tf


def bg(slide):
    rect(slide, 0, 0, W, H, BG)


def footer(slide, page, section=""):
    rect(slide, 0, Inches(7.28), W, Inches(0.22), RGBColor(0x08, 0x0D, 0x16))
    put(
        slide, Inches(0.4), Inches(7.28), Inches(9.2), Inches(0.22),
        ("TimesFM-3  ·  " + section) if section else "TimesFM-3 主要想法",
        9, MUTED, valign=MSO_ANCHOR.MIDDLE,
    )
    put(
        slide, Inches(11.3), Inches(7.28), Inches(1.6), Inches(0.22),
        f"{page}  /  {TOTAL}", 9, MUTED, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE,
    )


def header(slide, kicker, title, subtitle=None):
    rect(slide, 0, 0, Inches(0.12), H, CYAN)
    put(slide, Inches(0.45), Inches(0.14), Inches(12.4), Inches(0.26), kicker, 11, CYAN, True)
    put(slide, Inches(0.45), Inches(0.38), Inches(12.5), Inches(0.52), title, 22, WHITE, True)
    if subtitle:
        put(slide, Inches(0.45), Inches(0.88), Inches(12.5), Inches(0.34), subtitle, 13, MUTED)


def so_what(slide, text):
    box(slide, Inches(0.4), Inches(6.48), Inches(12.5), Inches(0.72), CARD2, AMBER)
    rect(slide, Inches(0.4), Inches(6.48), Inches(0.1), Inches(0.72), AMBER)
    put(slide, Inches(0.7), Inches(6.48), Inches(12.0), Inches(0.72),
        "要点  ·  " + text, 13, WHITE, True, valign=MSO_ANCHOR.MIDDLE)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # ========== 1 封面 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    rect(s, 0, 0, Inches(0.16), H, CYAN)
    put(s, Inches(0.7), Inches(1.2), Inches(12), Inches(0.32),
        "论文解读  ·  Google Research  ·  2026-08-31", 14, AMBER, True)
    put(s, Inches(0.7), Inches(1.65), Inches(12), Inches(0.75), "TimesFM-3 主要想法", 40, WHITE, True)
    put(s, Inches(0.7), Inches(2.45), Inches(12), Inches(0.4),
        "多变量时序预测：零样本、协变量、一次前向出整段", 18, SOFT)
    box(s, Inches(0.7), Inches(3.05), Inches(11.9), Inches(1.35), CARD, STROKE)
    put(s, Inches(0.95), Inches(3.2), Inches(11.4), Inches(1.1),
        "TimesFM 第三代：从零样本单变量，到原生零样本多变量。\n"
        "多目标 + 历史协变量 + 未来已知协变量，单次前向输出整段 horizon 的 9 个分位数。",
        16, SOFT)
    for i, (k, v) in enumerate([
        ("330M", "参数，20 层 transformer"),
        ("1T+", "预训练时间点（真实 + 合成）"),
        ("3 个第一", "GIFT-Eval / FEV-Bench / TIME"),
    ]):
        left = Inches(0.7 + i * 4.05)
        box(s, left, Inches(4.65), Inches(3.85), Inches(1.7), CARD2, STROKE)
        put(s, left + Inches(0.2), Inches(4.8), Inches(3.45), Inches(0.5), k, 22, CYAN, True)
        put(s, left + Inches(0.2), Inches(5.45), Inches(3.45), Inches(0.8), v, 14, WHITE)
    put(s, Inches(0.7), Inches(6.65), Inches(8), Inches(0.3),
        "来源：research.google/blog/timesfm-3  ·  本地存档 TimesFM-3-文章.md", 13, MUTED)
    put(s, Inches(10.3), Inches(6.65), Inches(2.5), Inches(0.3), f"1  /  {TOTAL}", 13, MUTED, align=PP_ALIGN.RIGHT)

    # ========== 2 一页看懂 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "0  一页看懂", "三句话讲完 TimesFM-3",
           "是什么、解决什么、关键招。")
    cards = [
        (CYAN, "是什么",
         "时序基础模型第三代。\n\n330M 参数、1T+ 时间点预训练，\n零样本直接预测新序列。",
         "前代只能单变量；3 代原生多变量。"),
        (AMBER, "解决什么",
         "真实预测本质是多变量：\n销量受相关产品、客流、\n促销计划共同影响。",
         "单变量外推看不见「未来已知事件」。"),
        (GREEN, "关键招",
         "交替注意力（时间 × 变量）+\n连续 patch 掩码：\n整段 horizon 一次前向填出。",
         "不再逐段自回归，误差不累积。"),
    ]
    for i, (c, t, a, b) in enumerate(cards):
        left = Inches(0.4 + i * 4.3)
        box(s, left, Inches(1.35), Inches(4.15), Inches(4.95), CARD, STROKE)
        rect(s, left, Inches(1.35), Inches(0.12), Inches(4.95), c)
        put(s, left + Inches(0.3), Inches(1.55), Inches(3.65), Inches(0.45), t, 18, c, True)
        put(s, left + Inches(0.3), Inches(2.2), Inches(3.65), Inches(2.4), a, 16, WHITE)
        put(s, left + Inches(0.3), Inches(4.7), Inches(3.65), Inches(1.3), b, 14, SOFT)
    footer(s, 2, "一页看懂")

    # ========== 3 为什么需要多变量 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "1  问题", "单变量外推，看不见「未来已知的事」",
           "例：预测冰淇淋销量，历史销量之外还有相关产品、客流、促销计划。")
    box(s, Inches(0.4), Inches(1.35), Inches(6.0), Inches(4.4), CARD, STROKE)
    put(s, Inches(0.65), Inches(1.55), Inches(5.5), Inches(0.4), "单变量模式（红线）", 17, MUTED, True)
    put(s, Inches(0.65), Inches(2.15), Inches(5.5), Inches(3.2),
        "只看历史销量，把周模式向前平移。\n\n"
        "不知道哪天有促销 → 促销日销量被系统性低估。\n\n"
        "前代 TimesFM（≤2.5）只能这样做。",
        15, SOFT)
    box(s, Inches(6.6), Inches(1.35), Inches(6.3), Inches(4.4), CARD2, CYAN)
    put(s, Inches(6.85), Inches(1.55), Inches(5.8), Inches(0.4), "TimesFM-3 多变量模式（蓝线）", 17, CYAN, True)
    put(s, Inches(6.85), Inches(2.15), Inches(5.8), Inches(3.2),
        "把「促销计划」作为未来已知协变量输入。\n\n"
        "从历史中学到促销 → 销量提升的关系，\n应用到未来促销日：每个促销日抬高约 20%。\n\n"
        "相关产品销量、历史客流也可一并输入。",
        15, WHITE)
    so_what(s, "预测质量的差距，主要来自「能不能用上其他序列和已知未来」，而不是更大的单序列模型。")
    footer(s, 3, "问题")

    # ========== 4 输入能力 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "2  输入能力", "三类序列，统一进同一个模型",
           "全部零样本：不需要针对任务微调。")
    rows = [
        ("多目标", "同时预测多条相关序列", "不同品牌冰淇淋的销量", "点预测 + 分位数预测", CYAN),
        ("历史协变量", "只有过去已知的特征", "历史客流、历史温度", "horizon 段被掩码", INDIGO),
        ("历史-未来协变量", "未来也已知的特征", "促销计划、节假日、天气预报", "horizon 段保持可见", AMBER),
    ]
    box(s, Inches(0.4), Inches(1.32), Inches(12.5), Inches(0.5), RGBColor(0x1A, 0x2A, 0x42))
    for j, h in enumerate(["类型", "含义", "例子", "解码时"]):
        put(s, Inches(0.6 + j * 3.1), Inches(1.32), Inches(2.9), Inches(0.5), h, 14, CYAN, True, valign=MSO_ANCHOR.MIDDLE)
    for i, (a, b, c, d, color) in enumerate(rows):
        top = Inches(1.92 + i * 1.35)
        box(s, Inches(0.4), top, Inches(12.5), Inches(1.2), CARD if i % 2 == 0 else CARD2, STROKE)
        rect(s, Inches(0.4), top, Inches(0.1), Inches(1.2), color)
        put(s, Inches(0.65), top, Inches(2.9), Inches(1.2), a, 16, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
        put(s, Inches(3.7), top, Inches(2.9), Inches(1.2), b, 15, SOFT, valign=MSO_ANCHOR.MIDDLE)
        put(s, Inches(6.8), top, Inches(2.9), Inches(1.2), c, 15, SOFT, valign=MSO_ANCHOR.MIDDLE)
        put(s, Inches(9.9), top, Inches(2.9), Inches(1.2), d, 15, WHITE, valign=MSO_ANCHOR.MIDDLE)
    so_what(s, "「未来已知协变量」是业务预测最常用、而单变量模型结构性用不上的信息。")
    footer(s, 4, "输入能力")

    # ========== 5 架构 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "3  架构", "patch 化 + 交替注意力：时间轴与变量轴分开做",
           "decoder-only transformer，20 层，model dim 1280，16 头。")
    cards = [
        ("Patch = 32 步", "连续 32 个时间点为一个 token；\n每条序列单独归一化，\n兼容量纲差异巨大的序列。", CYAN),
        ("多变量 token", "目标 / 历史协变量：单 patch 成 token。\n历史-未来协变量：拼接当前与未来 patch，\n让模型「看到」已知未来信号。", INDIGO),
        ("交替注意力", "因果时间注意力：沿时间轴，只看过去。\n全变量注意力：沿变量轴，看所有序列。\n两种交替堆叠，融合时间与跨序列关系。", AMBER),
    ]
    for i, (t, d, c) in enumerate(cards):
        left = Inches(0.4 + i * 4.3)
        box(s, left, Inches(1.35), Inches(4.15), Inches(4.4), CARD, STROKE)
        rect(s, left, Inches(1.35), Inches(4.15), Inches(0.1), c)
        put(s, left + Inches(0.25), Inches(1.6), Inches(3.7), Inches(0.45), t, 18, c, True)
        put(s, left + Inches(0.25), Inches(2.25), Inches(3.7), Inches(3.2), d, 15, WHITE)
    so_what(s, "把「时间因果」和「跨序列相关」拆成两种注意力交替做，是多变量零样本成立的关键设计。")
    footer(s, 5, "架构")

    # ========== 6 单次前向解码 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "4  解码", "连续 patch 掩码：整段 horizon 一次前向填出",
           "Contiguous Patch Masking，告别逐段自回归。")
    box(s, Inches(0.4), Inches(1.35), Inches(6.0), Inches(4.4), CARD, STROKE)
    put(s, Inches(0.65), Inches(1.55), Inches(5.5), Inches(0.4), "前代：自回归逐段生成", 17, MUTED, True)
    put(s, Inches(0.65), Inches(2.15), Inches(5.5), Inches(3.2),
        "一段一段往后生成。\n\n"
        "延迟高、算力贵。\n\n"
        "前一段的误差带进后一段，越往后越漂。",
        15, SOFT)
    box(s, Inches(6.6), Inches(1.35), Inches(6.3), Inches(4.4), CARD2, CYAN)
    put(s, Inches(6.85), Inches(1.55), Inches(5.8), Inches(0.4), "TimesFM-3：掩码一次填出", 17, CYAN, True)
    put(s, Inches(6.85), Inches(2.15), Inches(5.8), Inches(3.2),
        "horizon 段先放掩码占位 token：\n目标与历史协变量遮住，未来已知协变量可见。\n\n"
        "交替注意力同时填出所有掩码 patch，无迭代。\n\n"
        "每步输出 9 个分位数（10%–90%），自带不确定性。",
        15, WHITE)
    so_what(s, "一次前向 = 更快、更稳；9 分位数输出 = 直接可做区间预测与风险决策。")
    footer(s, 6, "解码")

    # ========== 7 效果 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "5  效果", "三个公开基准，点预测与概率预测双第一",
           "对比 Chronos-2、Toto 2.0 系列、TimesFM-2.5 等预训练基础模型。")
    rows = [
        ("GIFT-Eval", "基础模型中排名第一", "点预测 + 概率预测"),
        ("FEV-Bench", "100 个真实任务总排名第一", "点预测 + 概率预测"),
        ("TIME", "50 领域 / 98 任务总排名第一", "点预测 + 概率预测"),
    ]
    for i, (a, b, c) in enumerate(rows):
        top = Inches(1.35 + i * 1.15)
        box(s, Inches(0.4), top, Inches(12.5), Inches(1.0), CARD if i % 2 == 0 else CARD2, STROKE)
        put(s, Inches(0.65), top, Inches(3.2), Inches(1.0), a, 18, CYAN, True, valign=MSO_ANCHOR.MIDDLE)
        put(s, Inches(4.1), top, Inches(5.2), Inches(1.0), b, 16, WHITE, valign=MSO_ANCHOR.MIDDLE)
        put(s, Inches(9.5), top, Inches(3.2), Inches(1.0), c, 14, SOFT, valign=MSO_ANCHOR.MIDDLE)
    box(s, Inches(0.4), Inches(4.95), Inches(12.5), Inches(1.25), CARD, STROKE)
    put(s, Inches(0.65), Inches(5.1), Inches(12.0), Inches(1.0),
        "两个模式都要看：单变量模式（不用协变量）已匹配或超过其他可复现模型；多变量模式再上一个台阶。\n"
        "说明增益既来自更强的主干，也来自「用上跨序列与协变量」。",
        15, WHITE)
    footer(s, 7, "效果")

    # ========== 8 可用性与启示 ==========
    s = prs.slides.add_slide(blank)
    bg(s)
    header(s, "6  可用性与启示", "能拿什么，不能拿什么",
           "代码 Apache-2.0；3.0 权重为非商用许可（2.5 及之前为 Apache-2.0）。")
    box(s, Inches(0.4), Inches(1.35), Inches(6.0), Inches(4.4), CARD, STROKE)
    put(s, Inches(0.65), Inches(1.55), Inches(5.5), Inches(0.4), "可以直接用", 17, GREEN, True)
    put(s, Inches(0.65), Inches(2.15), Inches(5.5), Inches(3.2),
        "GitHub：google-research/timesfm。\n"
        "HF 权重：google/timesfm-3.0-pytorch。\n"
        "BigQuery AI.FORECAST 陆续上线。\n\n"
        "研究、评测、做基线：没问题。",
        15, WHITE)
    box(s, Inches(6.6), Inches(1.35), Inches(6.3), Inches(4.4), CARD2, AMBER)
    put(s, Inches(6.85), Inches(1.55), Inches(5.8), Inches(0.4), "对我们的启示", 17, AMBER, True)
    put(s, Inches(6.85), Inches(2.15), Inches(5.8), Inches(3.2),
        "抄思路：patch、decoder-only、零样本、\n概率预测、协变量接口、掩码式一次解码。\n\n"
        "不抄路径：权重非商用、云预测为主。\n"
        "物理传感器场景仍需自有 Ingest / Tokenizer / Eval。",
        15, WHITE)
    so_what(s, "TimesFM-3 证明「时序基础模型 + 协变量接口」成立；它是我们评测里的通用 TSFM 基线。")
    footer(s, 8, "启示")

    out = r"g:\My Drive\Documents\AI物理感知\TimesFM-3-主要想法.pptx"
    try:
        prs.save(out)
    except PermissionError:
        out = r"g:\My Drive\Documents\AI物理感知\TimesFM-3-主要想法-v2.pptx"
        prs.save(out)
    print("SAVED", out)
    print("SLIDES", len(prs.slides))


if __name__ == "__main__":
    build()
