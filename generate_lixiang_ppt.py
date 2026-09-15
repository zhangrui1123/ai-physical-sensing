# -*- coding: utf-8 -*-
"""AI 感知立项 PPT：RSI 三要素 + 场景实例化 + 路标互锁。2 页。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

TEMPLATE = r"g:\My Drive\Documents\AI物理感知\PPT模板-浅色版16-9.pptx"
OUT = r"g:\My Drive\Documents\AI物理感知\AI感知立项-RSI循环.pptx"

BG = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF7, 0xF7, 0xF7)
STROKE = RGBColor(0xDD, 0xDD, 0xDD)
RED = RGBColor(0xC7, 0x00, 0x0B)
NAVY = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x55, 0x57, 0x57)
MUTED = RGBColor(0x89, 0x89, 0x89)
CYAN = RGBColor(0x30, 0xB5, 0xC5)
GREEN = RGBColor(0x62, 0xB2, 0x30)
ORANGE = RGBColor(0xED, 0x6D, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINK = RGBColor(0xFF, 0xF0, 0xF0)
FONT = "Microsoft YaHei"
TOTAL = 2


def set_run(run, size, color, bold=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25) if line == RED else Pt(0.75)
    try:
        sh.adjustments[0] = 0.05
    except Exception:
        pass
    return sh


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def put(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, after=2):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        run = p.add_run()
        run.text = line
        set_run(run, size, color, bold)
    return tf


def delete_all_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def blank_slide(prs):
    layout = min(prs.slide_layouts, key=lambda L: len(L.placeholders))
    s = prs.slides.add_slide(layout)
    for sh in list(s.placeholders):
        sp = sh._element
        sp.getparent().remove(sp)
    rect(s, 0, 0, prs.slide_width, prs.slide_height, BG)
    return s


def footer(slide, page, section=""):
    rect(slide, 0, Inches(7.15), Inches(13.333), Inches(0.35), RGBColor(0xF2, 0xF2, 0xF2))
    put(
        slide, Inches(0.5), Inches(7.15), Inches(9.0), Inches(0.35),
        f"华为  ·  AI 感知立项  ·  {section}    |    Security Level: Internal",
        9, MUTED, valign=MSO_ANCHOR.MIDDLE,
    )
    put(
        slide, Inches(11.2), Inches(7.15), Inches(1.7), Inches(0.35),
        f"{page}  /  {TOTAL}", 9, MUTED, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE,
    )


def header(slide, kicker, title):
    put(slide, Inches(0.5), Inches(0.18), Inches(12.2), Inches(0.26), kicker, 11, RED, True)
    put(slide, Inches(0.5), Inches(0.42), Inches(12.2), Inches(0.42), title, 22, NAVY, True)
    rect(slide, Inches(0.5), Inches(0.88), Inches(1.2), Inches(0.05), RED)


def so_what(slide, text):
    box(slide, Inches(0.5), Inches(6.32), Inches(12.3), Inches(0.74), PINK, RED)
    rect(slide, Inches(0.5), Inches(6.32), Inches(0.08), Inches(0.74), RED)
    put(slide, Inches(0.75), Inches(6.32), Inches(11.9), Inches(0.74),
        "所以  ·  " + text, 13, NAVY, True, valign=MSO_ANCHOR.MIDDLE)


def build():
    prs = Presentation(TEMPLATE)
    delete_all_slides(prs)

    # ========== 1  RSI 抽象 + 握姿实例化 ==========
    s = blank_slide(prs)
    header(s, "立项  ·  AI 感知研究方向  ·  北极星：关键场景感知精度 > 99%",
           "RSI 三要素：改什么、谁来判、分数怎么回到下一轮")

    loops = [
        ("小循环  ·  一次训练", GREEN, False,
         "改什么（决策变量）\n学习率、增强、阈值、早停。",
         "谁来判（评估器）\n验证集 loss、单次四分类、200ms 内是否稳住。",
         "分数怎么回来\n一次训练结束打分，留下更好的超参再训。",
         "上期人手    本期仍人手"),
        ("中循环  ·  算法方案  ·  本期自动", RED, True,
         "改什么（决策变量）\n融合怎么拼、模型骨架、训练配方、窗口与拒识门限。",
         "谁来判（评估器）\n冻结评测集：精度、跨机型衰减、时延、误触。",
         "分数怎么回来\n出方案 → 打分 → 更好才留下 → 自动再出一版。",
         "上期人手    本期机器转"),
        ("大循环  ·  问题定义", ORANGE, False,
         "改什么（决策变量）\n关键场景边界、>99% 口径、类别够不够、数据补哪里。",
         "谁来判（评估器）\n现网跟手率、误触、机型覆盖、失效清单是否关掉。",
         "分数怎么回来\n只回答「问题要不要改」。算法不许改自己的评分规则。",
         "上期人手    本期仍人拍板"),
    ]
    for i, (title, color, hot, a, b, c, who) in enumerate(loops):
        left = Inches(0.45 + i * 4.25)
        fill = PINK if hot else WHITE
        box(s, left, Inches(1.08), Inches(4.1), Inches(3.18), fill, color)
        rect(s, left, Inches(1.08), Inches(4.1), Inches(0.08), color)
        put(s, left + Inches(0.14), Inches(1.2), Inches(3.82), Inches(0.36), title, 14, color, True)
        put(s, left + Inches(0.14), Inches(1.58), Inches(3.82), Inches(0.62), a, 11, BODY, after=1)
        put(s, left + Inches(0.14), Inches(2.22), Inches(3.82), Inches(0.62), b, 11, BODY, after=1)
        put(s, left + Inches(0.14), Inches(2.86), Inches(3.82), Inches(0.78), c, 11, BODY, after=1)
        put(s, left + Inches(0.14), Inches(3.72), Inches(3.82), Inches(0.4), who, 12, color, True)

    box(s, Inches(0.45), Inches(4.38), Inches(12.4), Inches(1.82), CARD, STROKE)
    put(s, Inches(0.65), Inches(4.46), Inches(12.0), Inches(0.3),
        "落到智感握姿：边框电容 + IMU + 接近光  →  左手 / 右手 / 双手 / 未握持", 13, RED, True)
    put(s, Inches(0.65), Inches(4.8), Inches(12.0), Inches(1.28),
        "小循环：调电容阈值、IMU 增强、切换平滑。裁判 = 验证集四分类。\n"
        "中循环：启发式拼接 vs 统一 token 融合；要不要线性探针；壳 / 手套怎么进训练集。"
        "裁判 = 冻结集精度、跨机型掉分、切换 ≤ 200ms。\n"
        "大循环：壳 > 3mm、戴手套算不算关键场景；>99% 含不含切换过程。口径一改，评估器重写，中循环重跑。\n"
        "上期：三层都是人改、人跑、人看。本期：只让中层自动改、自动跑、自动留。",
        12, BODY, after=3)
    so_what(s, "上期证明三层循环都能用手转；精度从「能用」到「>99%」卡在算法方案搜索。本期只自动转中循环。")
    footer(s, 1, "RSI 抽象与实例化")

    # ========== 2  技术路标 × 应用路标 ==========
    s = blank_slide(prs)
    header(s, "立项  ·  2～3 年路标",
           "应用定「评到哪」，技术定「谁来改」，两条路标锁在一起")

    years = [
        ("Y1  本期  ·  2026–27", RED,
         "应用\n智感握姿 + 手表手势（敲一敲 / 划一划），关键场景精度 > 99%，现网可交。",
         "技术\n中循环自动迭代上线：冻结评估器 + 方案搜索器。小循环、大循环仍人手。",
         "互锁\n握姿 / 手势的精度、时延、误触写成中循环评估器。搜不到 >99%，应用升不了级。"),
        ("Y2  +1 年  ·  2027–28", ORANGE,
         "应用\n换机、换人、跑步 / 骑行 / 滑雪，精度不掉档；左右手、壳、手套也能用。",
         "技术\n评估器加上跨设备 / 跨用户衰减；个性化适配纳入中循环搜索空间。",
         "互锁\n应用要「换机换人还 >99%」，技术就把掉分写进裁判。不写，机器会刷单机分数。"),
        ("Y3  +2 年  ·  2028–29", CYAN,
         "应用\n新特性注册即用，不再为每个场景重训；延展到运动健康 / 工业。",
         "技术\n大循环半自动：机器建议场景边界，人拍板。物理 token 进入大模型。",
         "互锁\n应用要「新场景不重训」，评估器必须可插拔、接口不变。大循环仍不让算法改口径。"),
    ]
    for i, (title, color, app, tech, lock) in enumerate(years):
        left = Inches(0.45 + i * 4.25)
        box(s, left, Inches(1.08), Inches(4.1), Inches(3.92), WHITE, color)
        rect(s, left, Inches(1.08), Inches(4.1), Inches(0.08), color)
        put(s, left + Inches(0.14), Inches(1.2), Inches(3.82), Inches(0.34), title, 14, color, True)
        put(s, left + Inches(0.14), Inches(1.58), Inches(3.82), Inches(0.95), app, 12, BODY, after=2)
        put(s, left + Inches(0.14), Inches(2.56), Inches(3.82), Inches(1.05), tech, 12, BODY, after=2)
        put(s, left + Inches(0.14), Inches(3.64), Inches(3.82), Inches(1.2), lock, 12, NAVY, True, after=2)

    box(s, Inches(0.45), Inches(5.12), Inches(12.4), Inches(1.08), CARD, STROKE)
    put(s, Inches(0.65), Inches(5.2), Inches(12.0), Inches(0.28),
        "互锁怎么转（一条环，不是两条平行线）", 13, RED, True)
    put(s, Inches(0.65), Inches(5.5), Inches(12.0), Inches(0.62),
        "应用给出场景和数字 → 写成评估器（冻结，循环外）→ 卡住中循环搜索 → 搜出的方案回填应用 → 应用加新场景 → 评估器升级。\n"
        "评估器不让被搜的算法改掉，防止刷分。>99% 只认冻结协议上的数。",
        13, BODY, after=2)
    so_what(s, "本期立项交付物 = 中循环自动迭代器 + 握姿/手势冻结评估器，用来锁死 Y1 的 >99%。")
    footer(s, 2, "路标互锁")

    try:
        prs.save(OUT)
        out = OUT
    except PermissionError:
        out = OUT.replace(".pptx", "-v2.pptx")
        prs.save(out)
    print("SAVED", out)
    print("SLIDES", len(prs.slides))


if __name__ == "__main__":
    build()
