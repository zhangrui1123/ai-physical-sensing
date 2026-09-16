# -*- coding: utf-8 -*-
"""AI 物理感知平台 PPT：基于华为浅色 16:9 模板母版。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

TEMPLATE = r"g:\My Drive\Documents\AI物理感知\PPT模板-浅色版16-9.pptx"
OUT = r"g:\My Drive\Documents\AI物理感知\AI物理感知平台-华为汇报.pptx"

BG = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xF7, 0xF7, 0xF7)
STROKE = RGBColor(0xDD, 0xDD, 0xDD)
RED = RGBColor(0xC7, 0x00, 0x0B)
NAVY = RGBColor(0x1A, 0x1A, 0x1A)
BODY = RGBColor(0x55, 0x57, 0x57)
MUTED = RGBColor(0x89, 0x89, 0x89)
CYAN = RGBColor(0x30, 0xB5, 0xC5)
GREEN = RGBColor(0x62, 0xB2, 0x30)
ORANGE = RGBColor(0xED, 0x6D, 0x00)
INDIGO = RGBColor(0x00, 0x74, 0xCC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Microsoft YaHei"
TOTAL = 12


def set_run(run, size, color, bold=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", name)


def box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    try:
        sh.adjustments[0] = 0.05
    except Exception:
        pass
    return sh


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def put(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        set_run(run, size, color, bold)
    return tf


def delete_all_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def blank_slide(prs):
    layout = min(prs.slide_layouts, key=lambda L: len(L.placeholders))
    s = prs.slides.add_slide(layout)
    for sh in list(s.placeholders):
        sp = sh._element
        sp.getparent().remove(sp)
    rect(s, 0, 0, prs.slide_width, prs.slide_height, BG)
    return s


def footer(slide, page, section=""):
    rect(slide, 0, Inches(7.15), Inches(13.333), Inches(0.35), RGBColor(0xF2, 0xF2, 0xF2))
    put(
        slide, Inches(0.5), Inches(7.15), Inches(9.0), Inches(0.35),
        f"华为  ·  AI 物理感知平台  ·  {section}    |    Security Level: Internal",
        9, MUTED, valign=MSO_ANCHOR.MIDDLE,
    )
    put(
        slide, Inches(11.2), Inches(7.15), Inches(1.7), Inches(0.35),
        f"{page}  /  {TOTAL}", 9, MUTED, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE,
    )


def header(slide, kicker, title, subtitle=None):
    put(slide, Inches(0.5), Inches(0.22), Inches(12.2), Inches(0.28), kicker, 11, RED, True)
    put(slide, Inches(0.5), Inches(0.48), Inches(12.2), Inches(0.48), title, 24, NAVY, True)
    rect(slide, Inches(0.5), Inches(1.0), Inches(1.2), Inches(0.05), RED)
    if subtitle:
        put(slide, Inches(0.5), Inches(1.12), Inches(12.2), Inches(0.32), subtitle, 12, MUTED)


def so_what(slide, text):
    box(slide, Inches(0.5), Inches(6.35), Inches(12.3), Inches(0.7), RGBColor(0xFF, 0xF0, 0xF0), RED)
    rect(slide, Inches(0.5), Inches(6.35), Inches(0.08), Inches(0.7), RED)
    put(slide, Inches(0.75), Inches(6.35), Inches(11.9), Inches(0.7),
        "所以  ·  " + text, 13, NAVY, True, valign=MSO_ANCHOR.MIDDLE)


def build():
    prs = Presentation(TEMPLATE)
    delete_all_slides(prs)
    pink = RGBColor(0xFF, 0xF0, 0xF0)

    # 1 封面
    s = blank_slide(prs)
    rect(s, 0, 0, Inches(0.12), prs.slide_height, RED)
    put(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.3),
        "华为内部立项  ·  从特性独立开发到统一特征底座", 14, RED, True)
    put(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.65), "AI 物理感知平台", 40, NAVY, True)
    put(s, Inches(0.7), Inches(2.28), Inches(12), Inches(0.38),
        "时间序列输入  →  时空特征传感器  →  特征输出；评估时按特性训练预测头 / 分类头", 16, BODY)
    box(s, Inches(0.7), Inches(2.85), Inches(11.9), Inches(1.35), CARD, STROKE)
    put(s, Inches(0.95), Inches(2.98), Inches(11.4), Inches(1.15),
        "现状：握姿、手势、旋转各自开发，模型与特征不通用，新增特性需重新训练。\n"
        "前期：智感握姿等特性已完成单点验证，明确了「不通用」的成因。\n"
        "本期：平台只输出特征，不承接场景；特性效果通过训练对应预测头 / 分类头评估。",
        14, BODY)
    for i, (k, v) in enumerate([
        ("现状", "各特性独立开发，特征不通用"),
        ("前期", "智感握姿 / 手表手势完成单点验证"),
        ("平台", "时序输入 · 特征输出 · 头仅用于评估"),
    ]):
        left = Inches(0.7 + i * 4.05)
        box(s, left, Inches(4.45), Inches(3.85), Inches(1.55), WHITE, STROKE)
        rect(s, left, Inches(4.45), Inches(3.85), Inches(0.08), RED)
        put(s, left + Inches(0.2), Inches(4.62), Inches(3.45), Inches(0.35), k, 13, RED, True)
        put(s, left + Inches(0.2), Inches(5.05), Inches(3.45), Inches(0.75), v, 15, NAVY)
    put(s, Inches(0.7), Inches(6.25), Inches(8.5), Inches(0.28),
        "标杆特性：手机智感握姿、手表手势（敲一敲 / 划一划）  ·  12 页", 12, MUTED)
    put(s, Inches(10.3), Inches(6.25), Inches(2.5), Inches(0.28), f"1  /  {TOTAL}", 12, MUTED, align=PP_ALIGN.RIGHT)
    put(s, Inches(0.7), Inches(6.65), Inches(6), Inches(0.25), "Security Level: Internal", 10, MUTED)

    # 2 现状：特性独立开发
    s = blank_slide(prs)
    header(s, "1  现状", "各特性独立开发：模型与特征不通用，新增特性需重建",
           "问题不在单点算法，而在特征与特性绑定——换一个特性就要重做一套。")
    pains = [
        ("智感握姿  ·  独立链路", "电容阵列 + IMU + 接近光，启发式融合。\n四分类与切换，仅服务来电交互。\n特征无法复用，换机型即返回 801。", RED),
        ("手表手势  ·  独立链路", "IMU + X-TAP 单独建模，与握姿不共享。\nWATCH 3 与 WATCH 5 手势集不通用。\n跑步 / 骑行 / 滑雪场景下整体失效。", ORANGE),
        ("旋转 / 误触  ·  各自独立", "智感旋转基于陀螺仪，误触另设门限。\n同一只手表、同一路 IMU，特征不能复用。\n新增特性 = 新模型 + 新数据 + 新标定。", CYAN),
    ]
    for i, (t, d, c) in enumerate(pains):
        left = Inches(0.45 + i * 4.25)
        box(s, left, Inches(1.55), Inches(4.1), Inches(3.7), WHITE, STROKE)
        rect(s, left, Inches(1.55), Inches(4.1), Inches(0.08), c)
        put(s, left + Inches(0.16), Inches(1.78), Inches(3.78), Inches(0.5), t, 15, c, True)
        put(s, left + Inches(0.16), Inches(2.4), Inches(3.78), Inches(2.6), d, 13, BODY)
    box(s, Inches(0.5), Inches(5.45), Inches(12.3), Inches(0.75), CARD, STROKE)
    put(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.55),
        "共性：输入都是时间序列，输出却是绑死的标签。缺的不是再做一个特性模型，是中间那一层通用特征。",
        14, RED, True)
    footer(s, 2, "现状")

    # 3 前期工作
    s = blank_slide(prs)
    header(s, "2  前期工作", "智感握姿等特性已完成单点验证，并定位了「不通用」的成因",
           "前期工作的价值在于：单特性已可商用，同时明确了特征无法跨特性复用。")
    works = [
        ("智感握姿", RED,
         "电容阵列 + IMU + 接近光。\n左手 / 右手 / 双手 / 未握持，200ms 内输出；支持来电交互。\n现网：MultimodalAwarenessKit，覆盖部分机型。"),
        ("手表手势", ORANGE,
         "IMU + X-TAP（PPG / ECG / 触摸）。\n敲一敲 / 划一划 / 翻腕，支持接电话、遥控拍照、车钥匙。\n现网：WATCH 5 首配 NPU。"),
        ("定位的缺口", CYAN,
         "保护壳 > 3mm、戴手套时精度下降；换机型需重新标定。\n握姿特征无法用于手势，手势特征无法用于旋转。\n新增特性仍需从零搭建完整链路。"),
    ]
    for i, (t, c, d) in enumerate(works):
        left = Inches(0.45 + i * 4.25)
        box(s, left, Inches(1.55), Inches(4.1), Inches(3.7), WHITE, STROKE)
        rect(s, left, Inches(1.55), Inches(4.1), Inches(0.08), c)
        put(s, left + Inches(0.16), Inches(1.78), Inches(3.78), Inches(0.45), t, 16, c, True)
        put(s, left + Inches(0.16), Inches(2.35), Inches(3.78), Inches(2.7), d, 13, BODY)
    box(s, Inches(0.5), Inches(5.45), Inches(12.3), Inches(0.75), CARD, STROKE)
    put(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.55),
        "前期交付的是特性算法；本期要交付的是各特性可共用的中间层——时空特征传感器。",
        14, RED, True)
    footer(s, 3, "前期工作")

    # 4 平台定义
    s = blank_slide(prs)
    header(s, "3  平台定义", "平台只做一件事：把时间序列变成可复用的特征",
           "时空特征传感器 = 同时看「时间怎么变」和「多路传感器怎么摆」。")
    steps = [
        ("1  输入", "时间序列信号", "序列 id / 时间戳 / 数值\n量纲 / 质量位 / 变元角色\n电容 · IMU · 接近光 · X-TAP", RED),
        ("2  核心", "时空特征传感器", "时：窗口里信号怎么变\n空：多路传感器、阵列位置\n对齐时钟、量纲、质量位", ORANGE),
        ("3  输出", "特征", "同一套向量，不含场景标签\n握姿、手势、旋转都吃它\n底座可换，特征形态冻结", GREEN),
        ("4  仅评估", "预测头 / 分类头", "按特性另训一个小头\n头的分数 = 特征好不好\n头不是平台交付物", CYAN),
    ]
    for i, (k, t, d, c) in enumerate(steps):
        left = Inches(0.4 + i * 3.23)
        fill = pink if i == 1 else WHITE
        box(s, left, Inches(1.55), Inches(3.08), Inches(3.85), fill, c)
        rect(s, left, Inches(1.55), Inches(3.08), Inches(0.08), c)
        put(s, left + Inches(0.14), Inches(1.72), Inches(2.8), Inches(0.32), k, 12, c, True)
        put(s, left + Inches(0.14), Inches(2.08), Inches(2.8), Inches(0.45), t, 16, NAVY, True)
        put(s, left + Inches(0.14), Inches(2.6), Inches(2.8), Inches(2.5), d, 13, BODY)
    so_what(s, "平台边界：交付接入校验、时空特征传感器与特征。预测头 / 分类头仅在评测时训练，不进入平台接口。")
    footer(s, 4, "平台定义")

    # 5 输入规格
    s = blank_slide(prs)
    header(s, "4  输入", "入口只有一种东西：时间序列信号",
           "不管来自握姿、手势还是旋转，先变成同一套字段，缺一不可。")
    box(s, Inches(0.5), Inches(1.5), Inches(7.4), Inches(4.6), WHITE, STROKE)
    rect(s, Inches(0.5), Inches(1.5), Inches(7.4), Inches(0.08), RED)
    put(s, Inches(0.7), Inches(1.72), Inches(7.0), Inches(0.4), "六字段（缺一即拒收）", 17, RED, True)
    put(s, Inches(0.7), Inches(2.25), Inches(7.0), Inches(3.6),
        "序列 id：这一路信号是谁。\n"
        "时间戳：什么时候采到的——多路要对齐。\n"
        "数值：原始读数。\n"
        "量纲单位：g、rad/s、pF，不能混着喂。\n"
        "质量位：丢包、饱和、超档重采样，要记下来。\n"
        "变元角色：目标 / 仅历史 / 历史+未来。\n\n"
        "覆盖：IMU / 电容阵列 / 接近光 / X-TAP / EMG / TP。\n"
        "上限：变元 ≤ 32；历史段 ≤ 15,360 点。\n"
        "采样率档位：10 / 50 / 100 / 200 Hz。",
        14, BODY)
    box(s, Inches(8.1), Inches(1.5), Inches(4.7), Inches(4.6), CARD, STROKE)
    put(s, Inches(8.3), Inches(1.72), Inches(4.3), Inches(0.4), "不进入口的东西", 17, MUTED, True)
    put(s, Inches(8.3), Inches(2.25), Inches(4.3), Inches(3.6),
        "左手 / 右手 / 敲一敲\n——场景标签不进平台。\n\n"
        "来电交互、车钥匙解锁\n——应用语义不进平台。\n\n"
        "机型名、保护壳厚度\n——可作为质量位或协变量，不作为类别。",
        14, BODY)
    footer(s, 5, "输入")

    # 6 时空特征传感器
    s = blank_slide(prs)
    header(s, "5  时空特征传感器", "同时看时间怎么变、多路传感器怎么摆",
           "这是平台唯一的核心模块；换底座可以，进出形态不能散。")
    rows = [
        ("时  ·  时间维", "把一段窗口切成 patch，看信号怎么随时间变。\n握姿是秒级稳态，敲一敲是 0.5s 脉冲——同一套特征，窗口不同。", RED),
        ("空  ·  传感器维", "电容阵列有位置，IMU 有三轴，接近光 / X-TAP 是另一路。\n变元注意力自动学习多路信号的关联，不再人工拼接特征。", ORANGE),
        ("对齐先于模型", "时钟、量纲、质量位在进传感器之前完成。\n多路信号若未对齐，再强的模型也只能处理噪声。", CYAN),
    ]
    for i, (k, v, c) in enumerate(rows):
        top = Inches(1.5 + i * 1.35)
        box(s, Inches(0.5), top, Inches(12.3), Inches(1.22), WHITE, STROKE)
        rect(s, Inches(0.5), top, Inches(0.1), Inches(1.22), c)
        put(s, Inches(0.78), top + Inches(0.14), Inches(2.6), Inches(0.95), k, 16, c, True)
        put(s, Inches(3.5), top + Inches(0.16), Inches(9.05), Inches(0.95), v, 14, BODY)
    box(s, Inches(0.5), Inches(5.65), Inches(12.3), Inches(0.55), CARD, STROKE)
    put(s, Inches(0.7), Inches(5.72), Inches(11.9), Inches(0.42),
        "公开依据：TimesFM-3 变元注意力、UniTS 任务 token、LIMU-BERT-X 跨 1.1K 机型、Babel 6 模态对齐。",
        13, BODY)
    footer(s, 6, "特征传感器")

    # 7 输出：特征
    s = blank_slide(prs)
    header(s, "6  输出", "平台出口是特征，不是左手 / 敲一敲 / 转了多少度",
           "场景语义留在特性侧；特征形态冻结，换底座旧调用还能接。")
    box(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(4.55), WHITE, STROKE)
    rect(s, Inches(0.5), Inches(1.5), Inches(6.1), Inches(0.08), GREEN)
    put(s, Inches(0.7), Inches(1.72), Inches(5.7), Inches(0.4), "交什么（特征）", 17, GREEN, True)
    put(s, Inches(0.7), Inches(2.25), Inches(5.7), Inches(3.5),
        "一段窗口 → 一组向量（patch token）。\n\n"
        "同一套特征同时喂给：\n握姿分类头、手势分类头、旋转预测头。\n\n"
        "结构冻结：维度、时间对齐、质量位随路。\n版本升级向下兼容。\n\n"
        "底座（TimesFM / UniTS / 自研）可换，\n外面看到的仍是这组特征。",
        14, BODY)
    box(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(4.55), WHITE, STROKE)
    rect(s, Inches(6.8), Inches(1.5), Inches(6.0), Inches(0.08), MUTED)
    put(s, Inches(7.0), Inches(1.72), Inches(5.6), Inches(0.4), "不交什么", 17, MUTED, True)
    put(s, Inches(7.0), Inches(2.25), Inches(5.6), Inches(3.5),
        "不交付左手 / 右手标签。\n不交付敲一敲触发。\n不交付来电交互。\n不交付整机、芯片、场景 Pack。\n\n"
        "以上由特性团队负责。\n平台提供特征，特性团队训练各自的头。",
        14, BODY)
    footer(s, 7, "输出")

    # 8 评估：训练头
    s = blank_slide(prs)
    header(s, "7  评估", "特征冻结，按特性训练预测头 / 分类头，头的指标即特征质量",
           "预测头 / 分类头仅作评测探针，不作为平台交付物；头规模小，便于归因于特征。")
    box(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(1.55), WHITE, STROKE)
    rect(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.08), RED)
    put(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(1.2),
        "评测协议：冻结时空特征传感器（不调整或仅允许极小适配）→ 按特性训练线性探针 / 浅层头 → 在冻结集上评分。\n"
        "对照基线 = 前期手工算法（握姿启发式、手表现网手势）。同一数据下，头指标更优才说明特征更优。\n"
        "更换特性只更换头、不更换特征；若换头后指标显著下降，说明特征未通用，而非头训练不足。",
        14, BODY)
    heads = [
        ("分类头  ·  握姿", "四分类：左 / 右 / 双手 / 未握持\n看宏 F1、切换是否 ≤ 200ms", RED),
        ("分类头  ·  手势", "敲一敲 / 划一划 / 翻腕 + 拒识\n看召回 @ 误报、p95 ≤ 50ms", ORANGE),
        ("预测头  ·  旋转 / 切换", "下一时刻角度或握姿切换\n看 MAE、分位数覆盖率", CYAN),
    ]
    for i, (t, d, c) in enumerate(heads):
        left = Inches(0.5 + i * 4.2)
        box(s, left, Inches(3.25), Inches(4.0), Inches(1.85), WHITE, STROKE)
        rect(s, left, Inches(3.25), Inches(4.0), Inches(0.08), c)
        put(s, left + Inches(0.16), Inches(3.42), Inches(3.68), Inches(0.4), t, 15, c, True)
        put(s, left + Inches(0.16), Inches(3.88), Inches(3.68), Inches(1.05), d, 13, BODY)
    so_what(s, "平台验收不看「再交付一个握姿模型」，看「同一套特征更换头后，仍优于前期单点方案」。")
    footer(s, 8, "评估")

    # 9 标杆特性：前期 vs 本期
    s = blank_slide(prs)
    header(s, "8  标杆接入", "前期特性保留，改为「共用特征 + 各自的头」",
           "握姿、手势、旋转仍作为验收场景；变化在中间层，不在场景本身。")
    cols = [(0.5, 2.3), (2.85, 3.2), (6.15, 3.15), (9.4, 3.4)]
    for (cx, cw), htxt in zip(cols, ["特性", "前期（独立开发）", "本期（平台）", "评估头 / 指标"]):
        put(s, Inches(cx + 0.1), Inches(1.42), Inches(cw - 0.18), Inches(0.32), htxt, 12, MUTED, True)
    rows = [
        ("智感握姿", "电容+IMU+光启发式\n特征无法复用", "同一套时空特征\n秒级窗口", "分类头 · 四分类\n宏 F1；切换 ≤ 200ms"),
        ("敲一敲 / 划一划", "IMU+X-TAP 独立建模\n与握姿不共享", "同一套时空特征\n短脉冲窗口", "分类头 · 事件+拒识\n召回 @ 误报；p95 ≤ 50ms"),
        ("招一招 / 翻手腕", "每只手表单独标定", "换头不换特征", "分类头 · 事件\n跨机型衰减 ≤ 5pp"),
        ("智感旋转", "陀螺仪独立链路", "同一套特征用于预测", "预测头 · 连续量\n角度 MAE"),
        ("误触抑制", "各链路单独设门限", "基于特征的置信度", "拒识头 · 共用\n误触 < 1 次/天"),
    ]
    for i, row in enumerate(rows):
        top = Inches(1.8 + i * 0.78)
        box(s, Inches(0.5), top, Inches(12.3), Inches(0.7), WHITE if i % 2 == 0 else CARD, STROKE)
        for j, ((cx, cw), cell) in enumerate(zip(cols, row)):
            color = RED if j == 0 else (NAVY if j == 2 else BODY)
            put(s, Inches(cx + 0.1), top + Inches(0.04), Inches(cw - 0.18), Inches(0.62), cell, 12,
                color, j in (0, 2))
    footer(s, 9, "标杆")

    # 10 新特性怎么上
    s = blank_slide(prs)
    header(s, "9  新特性接入", "新特性 = 注册时序 + 训练一个头，不再重建完整链路",
           "时空特征传感器保持冻结；特性团队仅需确定窗口、类别表及预测头 / 分类头类型。")
    steps = [
        ("① 接入时序", "按六字段注册该路信号。\n电容、IMU、X-TAP 均可接入。", RED),
        ("② 输出特征", "复用已上线的时空特征传感器。\n不重新训练底座。", ORANGE),
        ("③ 训练头评估", "分类头或预测头，线性探针即可起步。\n指标达标方可准入。", GREEN),
        ("④ 特性侧组装", "窗口、类别表、拒识门限由特性侧定义。\n平台不承接场景语义。", CYAN),
    ]
    for i, (t, d, c) in enumerate(steps):
        left = Inches(0.45 + i * 3.2)
        box(s, left, Inches(1.55), Inches(3.05), Inches(3.55), WHITE, STROKE)
        rect(s, left, Inches(1.55), Inches(3.05), Inches(0.08), c)
        put(s, left + Inches(0.15), Inches(1.75), Inches(2.75), Inches(0.45), t, 16, c, True)
        put(s, left + Inches(0.15), Inches(2.3), Inches(2.75), Inches(2.5), d, 14, BODY)
    box(s, Inches(0.5), Inches(5.3), Inches(12.3), Inches(0.9), CARD, STROKE)
    put(s, Inches(0.7), Inches(5.42), Inches(11.9), Inches(0.7),
        "与前期差别：前期每新增一个特性，需重做模型、特征与标定；本期每新增一个特性，仅增加一个头。若头无法达标，回溯优化特征传感器，而非另立特性项目。",
        14, RED, True)
    footer(s, 10, "新特性")

    # 11 量化指标
    s = blank_slide(prs)
    header(s, "10  量化指标", "考核对象为特征，指标通过对应的头体现",
           "每项均为「冻结特征 + 对应的头」。基线 = 前期独立开发的手工算法。")
    cols = [(0.5, 1.3), (1.85, 2.3), (4.2, 2.6), (6.85, 3.3), (10.2, 2.6)]
    for (cx, cw), htxt in zip(cols, ["层", "指标", "基线 / 对标", "目标值", "怎么测"]):
        put(s, Inches(cx + 0.1), Inches(1.42), Inches(cw - 0.18), Inches(0.32), htxt, 12, MUTED, True)
    rows = [
        ("特征", "握姿分类头 宏 F1", "前期启发式握姿", "≥ 前期；关键场景冲 99%", "冻结特征 + 四分类头"),
        ("特征", "手势分类头 宏 F1", "WATCH 5 现网手势", "≥ 前期；误触 < 1 次/天", "冻结特征 + 事件头"),
        ("特征", "旋转预测头 MAE", "前期陀螺仪单链", "持平或更优", "冻结特征 + 预测头"),
        ("通用", "换头掉分", "各特性独立模型", "换头不换特征，掉分 ≤ 5pp", "握姿特征直接挂手势头"),
        ("泛化", "跨机型 / 跨用户", "单机单独标定", "衰减 ≤ 5pp；Gen2 ≤ 2pp", "留一交叉验证"),
        ("时延", "特征 + 头 端侧 p95", "现网握姿 200ms", "握姿 ≤ 200ms；手势 ≤ 50ms", "真机 profiling"),
        ("工程", "特征形态兼容", "—", "换底座不破坏旧调用", "工具链回归"),
    ]
    for i, row in enumerate(rows):
        top = Inches(1.8 + i * 0.62)
        box(s, Inches(0.5), top, Inches(12.3), Inches(0.56), WHITE if i % 2 == 0 else CARD, STROKE)
        for j, ((cx, cw), cell) in enumerate(zip(cols, row)):
            color = RED if j == 0 else (NAVY if j == 3 else BODY)
            put(s, Inches(cx + 0.1), top + Inches(0.03), Inches(cw - 0.18), Inches(0.5), cell, 10,
                color, j in (0, 3))
    box(s, Inches(0.5), Inches(6.25), Inches(12.3), Inches(0.75), CARD, STROKE)
    put(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.55),
        "验收标准：同一套特征分别训练握姿头与手势头，均优于前期单点方案——方为通用，而非再交付两个特性模型。",
        12, RED, True)
    footer(s, 11, "量化指标")

    # 12 代际演进
    s = blank_slide(prs)
    header(s, "11  代际演进", "特征先通用，再覆盖到人与设备，最终成为大模型可用的 token",
           "每代升级的是时空特征传感器；头可替换重训，特征形态保持稳定。")
    gens = [
        ("Gen1  特征通用",
         "时空特征传感器上线。\n握姿 / 手势通过线性探针头验收。\n优于前期单点，方为通用起步。\n回答：一套特征能否支撑多个特性。", RED),
        ("Gen2  跨设备跨人通用",
         "左右手、保护壳、手套、佩戴松紧纳入特征。\nLoRA 仅做特征侧轻量适配，不重建特性链路。\n跨机型衰减收紧至 ≤ 2pp。\n回答：换设备换人，是否仍需重训。", ORANGE),
        ("Gen3  特征接入大模型",
         "物理特征作为新模态注入大模型。\n特性头可进一步简化，甚至不再独立存在。\n平台从「提供特征」演进为「提供 token」。\n回答：物理感知如何进入统一智能。", CYAN),
    ]
    for i, (t, d, c) in enumerate(gens):
        left = Inches(0.45 + i * 4.25)
        box(s, left, Inches(1.55), Inches(4.1), Inches(3.7), WHITE, STROKE)
        rect(s, left, Inches(1.55), Inches(4.1), Inches(0.08), c)
        put(s, left + Inches(0.16), Inches(1.78), Inches(3.78), Inches(0.45), t, 16, c, True)
        put(s, left + Inches(0.16), Inches(2.35), Inches(3.78), Inches(2.7), d, 13, BODY)
    box(s, Inches(0.5), Inches(5.45), Inches(12.3), Inches(0.75), CARD, STROKE)
    put(s, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.55),
        "演进原则：输入时序规格、输出特征形态向下兼容；底座和评估头可独立替换。",
        14, RED, True)
    footer(s, 12, "代际演进")

    try:
        prs.save(OUT)
        out = OUT
    except PermissionError:
        out = OUT.replace(".pptx", "-v2.pptx")
        prs.save(out)
    print("SAVED", out)
    print("SLIDES", len(prs.slides))


if __name__ == "__main__":
    build()


